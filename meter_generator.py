import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from datetime import datetime, timedelta
import pandas as pd
import random
import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows
import os

class MeterDataGenerator:
    def __init__(self, root):
        self.root = root
        self.root.title("Генератор даних лічильників")
        self.root.geometry("600x500")
        self.root.configure(bg='#f0f0f0')
        
        # Стилі
        style = ttk.Style()
        style.theme_use('clam')
        
        self.create_widgets()
    
    def create_widgets(self):
        # Головний фрейм
        main_frame = ttk.Frame(self.root, padding="20")
        main_frame.grid(row=0, column=0, sticky=(tk.W, tk.E, tk.N, tk.S))
        
        # Заголовок
        title_label = ttk.Label(main_frame, text="Генератор даних лічильників", 
                               font=('Arial', 16, 'bold'))
        title_label.grid(row=0, column=0, columnspan=2, pady=(0, 20))
        
        # Дата початку
        ttk.Label(main_frame, text="Дата початку (РРРР-ММ-ДД):").grid(row=1, column=0, sticky=tk.W, pady=5)
        self.date_entry = ttk.Entry(main_frame, width=20)
        self.date_entry.insert(0, "2024-01-01")
        self.date_entry.grid(row=1, column=1, sticky=tk.W, pady=5)
        
        # Номер лічильника
        ttk.Label(main_frame, text="Номер лічильника:").grid(row=2, column=0, sticky=tk.W, pady=5)
        self.meter_number_entry = ttk.Entry(main_frame, width=20)
        self.meter_number_entry.insert(0, "001")
        self.meter_number_entry.grid(row=2, column=1, sticky=tk.W, pady=5)
        
        # Час початку
        time_frame = ttk.Frame(main_frame)
        time_frame.grid(row=3, column=0, columnspan=2, sticky=tk.W, pady=5)
        
        ttk.Label(time_frame, text="Час початку:").grid(row=0, column=0, sticky=tk.W)
        
        self.hour_var = tk.StringVar(value="00")
        self.minute_var = tk.StringVar(value="00")
        
        hour_spinbox = ttk.Spinbox(time_frame, from_=0, to=23, width=5, 
                                  textvariable=self.hour_var, format="%02.0f")
        hour_spinbox.grid(row=0, column=1, padx=(10, 5))
        
        ttk.Label(time_frame, text=":").grid(row=0, column=2)
        
        minute_spinbox = ttk.Spinbox(time_frame, values=[f"{i:02d}" for i in range(0, 60, 10)], 
                                   width=5, textvariable=self.minute_var)
        minute_spinbox.grid(row=0, column=3, padx=(5, 0))
        
        # Діапазон напруги
        ttk.Label(main_frame, text="Мінімальна напруга (В):").grid(row=4, column=0, sticky=tk.W, pady=5)
        self.min_voltage_entry = ttk.Entry(main_frame, width=20)
        self.min_voltage_entry.insert(0, "220.00")
        self.min_voltage_entry.grid(row=4, column=1, sticky=tk.W, pady=5)
        
        ttk.Label(main_frame, text="Максимальна напруга (В):").grid(row=5, column=0, sticky=tk.W, pady=5)
        self.max_voltage_entry = ttk.Entry(main_frame, width=20)
        self.max_voltage_entry.insert(0, "240.00")
        self.max_voltage_entry.grid(row=5, column=1, sticky=tk.W, pady=5)
        
        # Тип лічильника
        ttk.Label(main_frame, text="Тип лічильника:").grid(row=6, column=0, sticky=tk.W, pady=5)
        self.meter_type_var = tk.StringVar(value="1-фазний")
        meter_type_combo = ttk.Combobox(main_frame, textvariable=self.meter_type_var, 
                                       values=["1-фазний", "3-фазний"], state="readonly", width=17)
        meter_type_combo.grid(row=6, column=1, sticky=tk.W, pady=5)
        
        # Кнопки
        button_frame = ttk.Frame(main_frame)
        button_frame.grid(row=7, column=0, columnspan=2, pady=(30, 10))
        
        generate_btn = ttk.Button(button_frame, text="Генерувати дані", 
                                command=self.generate_data, style='Accent.TButton')
        generate_btn.pack(side=tk.LEFT, padx=(0, 10))
        
        save_btn = ttk.Button(button_frame, text="Зберегти Excel", 
                            command=self.save_excel)
        save_btn.pack(side=tk.LEFT)
        
        # Прогрес бар
        self.progress = ttk.Progressbar(main_frame, length=400, mode='determinate')
        self.progress.grid(row=8, column=0, columnspan=2, pady=10, sticky=(tk.W, tk.E))
        
        # Статус
        self.status_label = ttk.Label(main_frame, text="Готово до роботи", 
                                    foreground='green')
        self.status_label.grid(row=9, column=0, columnspan=2, pady=5)
        
        self.data = None
    
    def validate_input(self):
        try:
            # Перевірка дати
            datetime.strptime(self.date_entry.get(), "%Y-%m-%d")
            
            # Перевірка напруги
            min_voltage = float(self.min_voltage_entry.get())
            max_voltage = float(self.max_voltage_entry.get())
            
            if min_voltage >= max_voltage:
                raise ValueError("Мінімальна напруга повинна бути менше максимальної")
            
            # Перевірка номера лічильника
            if not self.meter_number_entry.get().strip():
                raise ValueError("Номер лічильника не може бути пустим")
            
            return True, min_voltage, max_voltage
            
        except ValueError as e:
            messagebox.showerror("Помилка", f"Некоректні дані: {str(e)}")
            return False, None, None
    
    def generate_data(self):
        valid, min_voltage, max_voltage = self.validate_input()
        if not valid:
            return
        
        self.status_label.config(text="Генерація даних...", foreground='blue')
        self.progress['value'] = 0
        self.root.update()
        
        try:
            # Початкові параметри
            start_date = datetime.strptime(self.date_entry.get(), "%Y-%m-%d")
            start_time = datetime.combine(start_date.date(), 
                                        datetime.strptime(f"{self.hour_var.get()}:{self.minute_var.get()}", 
                                                        "%H:%M").time())
            
            meter_number = self.meter_number_entry.get().strip()
            is_three_phase = self.meter_type_var.get() == "3-фазний"
            
            # Генерація даних
            data = []
            current_time = start_time
            
            for i in range(1200):
                row = {
                    'Номер лічильника': meter_number,
                    'Дата': current_time.strftime("%Y-%m-%d"),
                    'Час': current_time.strftime("%H:%M"),
                    'Фаза A': round(random.uniform(min_voltage, max_voltage), 2)
                }
                
                if is_three_phase:
                    row['Фаза B'] = round(random.uniform(min_voltage, max_voltage), 2)
                    row['Фаза C'] = round(random.uniform(min_voltage, max_voltage), 2)
                
                data.append(row)
                current_time += timedelta(minutes=10)
                
                # Оновлення прогресу
                if i % 50 == 0:
                    self.progress['value'] = (i / 1200) * 100
                    self.root.update()
            
            self.data = pd.DataFrame(data)
            self.progress['value'] = 100
            self.status_label.config(text=f"Згенеровано {len(data)} записів", foreground='green')
            
        except Exception as e:
            messagebox.showerror("Помилка", f"Помилка генерації даних: {str(e)}")
            self.status_label.config(text="Помилка генерації", foreground='red')
    
    def save_excel(self):
        if self.data is None:
            messagebox.showwarning("Попередження", "Спочатку згенеруйте дані")
            return
        
        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx")],
            title="Зберегти файл Excel"
        )
        
        if not file_path:
            return
        
        self.status_label.config(text="Збереження Excel файлу...", foreground='blue')
        self.progress['value'] = 0
        self.root.update()
        
        try:
            # Створення Excel файлу
            with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
                # Записуємо дані на перший лист
                self.data.to_excel(writer, sheet_name='Дані', index=False)
                
                # Отримуємо workbook та worksheet
                workbook = writer.book
                worksheet = writer.sheets['Дані']
                
                # Автоширина колонок
                for column in worksheet.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    worksheet.column_dimensions[column_letter].width = max_length + 2
                
                self.progress['value'] = 50
                self.root.update()
                
                # Створення діаграми
                self.create_chart(workbook)
                
            self.progress['value'] = 100
            self.status_label.config(text=f"Файл збережено: {file_path}", foreground='green')
            messagebox.showinfo("Успіх", f"Файл Excel збережено:\n{file_path}")
            
        except Exception as e:
            messagebox.showerror("Помилка", f"Помилка збереження файлу: {str(e)}")
            self.status_label.config(text="Помилка збереження", foreground='red')
    
    def create_chart(self, workbook):
        # Створення другого листа для діаграми
        chart_sheet = workbook.create_sheet(title="Діаграма")
        
        # Підготовка даних для діаграми (середні значення по годинах)
        self.data['DateTime'] = pd.to_datetime(self.data['Дата'] + ' ' + self.data['Час'])
        hourly_data = self.data.groupby(self.data['DateTime'].dt.floor('H')).agg({
            'Фаза A': ['min', 'max', 'mean']
        }).round(2)
        
        is_three_phase = 'Фаза B' in self.data.columns
        
        if is_three_phase:
            hourly_data_b = self.data.groupby(self.data['DateTime'].dt.floor('H')).agg({
                'Фаза B': ['min', 'max', 'mean']
            }).round(2)
            hourly_data_c = self.data.groupby(self.data['DateTime'].dt.floor('H')).agg({
                'Фаза C': ['min', 'max', 'mean']
            }).round(2)
        
        # Записуємо дані для діаграми
        chart_data = []
        for i, (timestamp, row) in enumerate(hourly_data.iterrows()):
            chart_row = {
                'Час': timestamp.strftime('%H:%M'),
                'Фаза A (мін)': row[('Фаза A', 'min')],
                'Фаза A (макс)': row[('Фаза A', 'max')],
                'Фаза A (сер)': row[('Фаза A', 'mean')]
            }
            
            if is_three_phase:
                chart_row.update({
                    'Фаза B (мін)': hourly_data_b.iloc[i][('Фаза B', 'min')],
                    'Фаза B (макс)': hourly_data_b.iloc[i][('Фаза B', 'max')],
                    'Фаза B (сер)': hourly_data_b.iloc[i][('Фаза B', 'mean')],
                    'Фаза C (мін)': hourly_data_c.iloc[i][('Фаза C', 'min')],
                    'Фаза C (макс)': hourly_data_c.iloc[i][('Фаза C', 'max')],
                    'Фаза C (сер)': hourly_data_c.iloc[i][('Фаза C', 'mean')]
                })
            
            chart_data.append(chart_row)
        
        # Записуємо дані на лист діаграми
        chart_df = pd.DataFrame(chart_data)
        for r_idx, row in enumerate(dataframe_to_rows(chart_df, index=False, header=True), 1):
            for c_idx, value in enumerate(row, 1):
                chart_sheet.cell(row=r_idx, column=c_idx, value=value)
        
        # Створення лінійної діаграми
        chart = LineChart()
        chart.title = "Аналіз напруги по годинах"
        chart.style = 10
        chart.x_axis.title = 'Час'
        chart.y_axis.title = 'Напруга (В)'
        chart.width = 20
        chart.height = 12
        
        # Дані для діаграми
        data_range = Reference(chart_sheet, min_col=2, min_row=1, 
                              max_col=len(chart_df.columns), max_row=len(chart_df) + 1)
        cats = Reference(chart_sheet, min_col=1, min_row=2, max_row=len(chart_df) + 1)
        
        chart.add_data(data_range, titles_from_data=True)
        chart.set_categories(cats)
        
        # Розміщення діаграми
        chart_sheet.add_chart(chart, "A10")

def main():
    root = tk.Tk()
    app = MeterDataGenerator(root)
    root.mainloop()

if __name__ == "__main__":
    main()